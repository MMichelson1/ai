#!/usr/bin/env python3
"""Beginner-friendly content revision of the Vibe Night deck (Oct 27, 2026).

Usage: python3 scripts/vn_deck_revise.py IN.pptx OUT.pptx   (run after vn_deck_fixes.py)

- "Agent" becomes "assistant" in plain words (an agent is really an assistant that does stuff for you)
- New slide after Assistant vs. App: free assistants to try tonight
- Assistant recipe: what you want, context, target
- App recipe: what and why, who uses it, how it looks, what good looks like
- Tools slide in plain words, no code-style text
"""
import copy, sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.opc.constants import RELATIONSHIP_TYPE as RT

src, out = sys.argv[1], sys.argv[2]
p = Presentation(src)
S = lambda n: p.slides[n - 1]
def shp(n, sid): return next(s for s in S(n).shapes if s.shape_id == sid)
def place(sh, l=None, t=None, w=None, h=None):
    if l is not None: sh.left = Inches(l)
    if t is not None: sh.top = Inches(t)
    if w is not None: sh.width = Inches(w)
    if h is not None: sh.height = Inches(h)
def remove(sh): sh._element.getparent().remove(sh._element)
def set_text(sh, txt):
    runs = sh.text_frame.paragraphs[0].runs
    runs[0].text = txt
    for r in runs[1:]: r.text = ""
    for extra in sh.text_frame.paragraphs[1:]:
        for r in extra.runs: r.text = ""

WHITE, ORANGE, CYAN = RGBColor(0xFF, 0xFF, 0xFF), RGBColor(0xF5, 0xA6, 0x23), RGBColor(0x33, 0xDD, 0xFF)
SOFT = RGBColor(0xC8, 0xB8, 0xF8)

def write_paras(sh, paras, size, space_after=6):
    tf = sh.text_frame; tf.word_wrap = True
    body = tf._txBody
    for extra in tf.paragraphs[1:]: body.remove(extra._p)
    first = tf.paragraphs[0]
    for r in list(first.runs): first._p.remove(r._r)
    for i, runs in enumerate(paras):
        para = first if i == 0 else tf.add_paragraph()
        para.space_after = Pt(space_after); para.space_before = Pt(0)
        for text, bold, color, italic, *sz in runs:
            r = para.add_run(); r.text = text
            f = r.font; f.size = Pt(sz[0] if sz else size); f.bold = bold; f.italic = italic; f.name = "Calibri"; f.color.rgb = color
def L(label, text, lc=ORANGE, size=None):
    runs = [(label + "  ", True, lc, False), (text, False, WHITE, False)]
    return [r + (size,) for r in runs] if size else runs

# ---- slide 9: rules
set_text(shp(9, 123), "Everything runs in your browser. Free accounts are fine.")
set_text(shp(9, 126), "⚡  Context Is Everything")
set_text(shp(9, 127), "Tell the AI what you want, why you want it, and what good looks like.")

# ---- slide 11: no dashes in the copy
set_text(shp(11, 162), "The AI does the hard part and shows you the result. You look at what it made, not the code.")
set_text(shp(11, 165), "“Make it darker.” “Add a date.” “That button doesn’t work. Fix it.”")

# ---- slide 13: tools in plain words, nothing that looks like code
set_text(shp(13, 177), "PICK ONE YOU LIKE")
set_text(shp(13, 178), "Four Tools to Vibe Code Tonight")
set_text(shp(13, 179), "All free in your browser. Pick one you know, or try them all.")
set_text(shp(13, 181), "Where to find it")
rows = [(185, "claude.ai", 186, "A great writer and builder. Best for bigger apps and detailed instructions."),
        (189, "chatgpt.com", 190, "Fast and familiar. A good all-around helper."),
        (193, "gemini.google.com", 194, "Works with your Google account. Good for quick builds."),
        (197, "perplexity.ai", 198, "Searches the web and shows its sources. Great for research.")]
for us, url, ds, desc in rows:
    write_paras(shp(13, us), [[(url, False, CYAN, False)]], size=14, space_after=0)
    set_text(shp(13, ds), desc)

# ---- slide 14: assistant vs app
set_text(shp(14, 205), "AN AGENT IS REALLY AN ASSISTANT")
set_text(shp(14, 206), "Assistant  vs.  App")
set_text(shp(14, 208), "\U0001F916  Assistant (agent)")
set_text(shp(14, 209), "•  Does stuff for you, on your behalf")
set_text(shp(14, 210), "•  Searches, compares, and fills things in")
set_text(shp(14, 211), "•  Figures out HOW to get it done")
set_text(shp(14, 212), "Example: Plan my trip to Chicago. Flights, hotel, food, and fun.")
set_text(shp(14, 214), "•  A tool that does one thing well")
set_text(shp(14, 215), "•  Works the same way every time")
set_text(shp(14, 216), "•  Has buttons, forms, and screens")
set_text(shp(14, 217), "Example: A tracker that shows my week at a glance.")
set_text(shp(14, 219), "Tonight you build both. First an assistant, then an app.")
for sid, t in [(209, 2.00), (210, 2.48), (211, 2.96), (214, 2.00), (215, 2.48), (216, 2.96)]:
    place(shp(14, sid), t=t)

# ---- slide 15: exercise 1 intro
set_text(shp(15, 228), "Open your favorite AI. Tell your assistant about the trip you want.")

# ---- slide 16: the assistant recipe
set_text(shp(16, 241), "THE ASSISTANT RECIPE:  WHAT  ·  CONTEXT  ·  TARGET")
set_text(shp(16, 242), "Tell Your Assistant What You Want")
write_paras(shp(16, 244), [
    L("1. What do you want it to do?", "Plan a 4-day trip for two, from Atlanta to Chicago, November 12 to 15."),
    L("2. Context:", "We've never been there. We love museums, art, and natural history. Suggest what to see, where to eat, and the best area to stay."),
    L("3. Target:", "Keep flights, hotel, and getting around under $1,200 for two, or as close as you can. Show me a day-by-day plan with prices."),
    L("4. Questions:", "Ask me any questions that would help refine this before you start."),
    [("Say it the way you'd say it to a friend.", False, SOFT, True)],
], size=16, space_after=10)
set_text(shp(16, 253), "\U0001F4A1  No need to say \u201cact like a travel agent.\u201d Just say what you want done.")
for r in shp(16, 253).text_frame.paragraphs[0].runs: r.font.size = Pt(15); r.font.italic = False

# ---- slide 17 and 20: "agent" -> "assistant"
for n in (17, 20):
    for sh in S(n).shapes:
        if not sh.has_text_frame: continue
        for para in sh.text_frame.paragraphs:
            for r in para.runs:
                if "agent" in r.text: r.text = r.text.replace("agent", "assistant")

# ---- slide 18: the app recipe
set_text(shp(18, 271), "Build an App")
set_text(shp(18, 272), "Context is the most important part. Give the AI the whole picture.")
for sid in (273, 276, 277, 279, 280, 281, 282, 278): remove(shp(18, sid))
recipe = shp(18, 274); place(recipe, l=0.5, t=1.92, w=9.0, h=2.45); recipe.text_frame.vertical_anchor = 1
write_paras(recipe, [
    L("What and why:", "What are you building, and what problem does it solve?", size=17),
    L("Who uses it:", "Just you? Your team? Your kids?", size=17),
    L("How it looks:", "iPhone, web, or both? Simple or detailed? Any colors you love?", size=17),
    L("What good looks like:", "When it works, what do you see? Describe the finished app.", size=17),
    L("Questions:", "Ask me questions along the way if you're unsure. Don't guess.", size=17),
], size=17, space_after=13)
foot = shp(18, 275); place(foot, l=0.5, t=4.58, w=7.6, h=0.4)
write_paras(foot, [[("Have an idea? Use it. No idea? The next slide has three ready-made ones.", False, SOFT, True)]], size=14, space_after=0)

# ---- slide 19: starter ideas in the app recipe
ideas = {294: "I want a simple iPhone app to log my workouts so I stay consistent. Just for me. Big buttons, calm blue colors. Good looks like: one tap to log, and I see my weekly streak.",
         298: "I want a web app to log my sales calls and follow-ups so nothing slips. For me and my team. Clean and simple, works on phone and laptop. Good looks like: my week on one screen.",
         302: "I want a word guessing game my family can play in the browser. Five-letter words, six guesses. Bright colors, easy for kids. Good looks like: we play a round in two minutes."}
for sid, text in ideas.items():
    write_paras(shp(19, sid), [[(text, False, WHITE, False)]], size=13, space_after=0)

# ---- new slide after 14: free assistants to try tonight (built from slide 13's layout)
src_slide = S(13)
new = p.slides.add_slide(src_slide.slide_layout)
for ph in list(new.placeholders): remove(ph)
bg = src_slide._element.find('.//{http://schemas.openxmlformats.org/presentationml/2006/main}bg')
if bg is not None:
    new._element.cSld.insert(0, copy.deepcopy(bg))
logo_part = None
for sh in src_slide.shapes:
    el = copy.deepcopy(sh._element)
    new.shapes._spTree.append(el)
    if sh.shape_type == 13:
        rid = sh._element.blipFill.blip.rEmbed
        logo_part = src_slide.part.related_part(rid)
        new_rid = new.part.relate_to(logo_part, RT.IMAGE)
        el.blipFill.blip.rEmbed = new_rid
N = len(p.slides)
set_text(shp(N, 177), "WHAT'S POSSIBLE TODAY")
set_text(shp(N, 178), "Free Assistants to Try Tonight")
set_text(shp(N, 179), "No subscription needed. Free plans limit how much they do, so start small.")
set_text(shp(N, 180), "Assistant"); set_text(shp(N, 181), "Where to find it"); set_text(shp(N, 182), "What it does for you")
assist = [("Muse", "muse.ai or the Muse app", "Meta's new assistant. Plans, fills in forms, and books."),
          ("Comet", "Free browser from Perplexity", "Reads the page you're on and clicks through tasks."),
          ("Gemini", "gemini.google.com", "Researches and builds your plan. Deep Research is free."),
          ("Manus", "manus.im", "Works through a bigger job, step by step.")]
for (nm, where, what), (a, b, c) in zip(assist, [(184, 185, 186), (188, 189, 190), (192, 193, 194), (196, 197, 198)]):
    set_text(shp(N, a), nm)
    write_paras(shp(N, b), [[(where, False, CYAN, False)]], size=13, space_after=0)
    set_text(shp(N, c), what)
# move it to sit right after slide 14
lst = p.slides._sldIdLst; ids = list(lst)
lst.remove(ids[-1]); lst.insert(14, ids[-1])

# ---- What's Next: room for four upcoming events, one line each
nx = next(sl for sl in p.slides if any(sh.has_text_frame and sh.text_frame.text.strip() == "What's Next" for sh in sl.shapes))
evbox = None
for sh in list(nx.shapes):
    if not sh.has_text_frame: continue
    t = sh.text_frame.text.strip()
    if t == "{{EV1_NAME}}": evbox = sh
    elif t in ("{{EV1_DETAILS}}", "{{EV2_NAME}}", "{{EV2_DETAILS}}"): remove(sh)
place(evbox, l=0.5, t=3.92, w=6.1, h=1.12); evbox.text_frame.vertical_anchor = 1
write_paras(evbox, [[("{{EV%d_NAME}}" % i, True, ORANGE, False, 12), ("{{EV%d_SEP}}{{EV%d_DETAILS}}" % (i, i), False, WHITE, False, 11)] for i in range(1, 5)], size=11, space_after=2)
for sh in nx.shapes:
    if sh.has_text_frame and sh.text_frame.text.strip() == "Thanks for being here tonight.": place(sh, t=5.18)
    elif sh.shape_id == 346: place(sh, t=5.12)

# ---- slide 5: live numbers from the AIC platform (read Oct 27, 2026)
# Global 511,511 members (+50,162 in 30 days), 313 chapters, 75 countries.
# Atlanta 4,869 members, 115 events since July 2025, 24 organizers and leads.
# South Central 17 chapters, 7,835 members. Carolinas 6 chapters, 560 members.
body = shp(5, 60)
write_paras(body, [
    [("The AI Collective is the world\u2019s largest grassroots AI community.", False, RGBColor(0, 0, 0), False)],
    [("The AI Collective Global", True, RGBColor(0, 0, 0), False), (" = 511,000+ members, 300+ chapters in 75 countries", False, RGBColor(0, 0, 0), False)],
    [("We added 50,000+ members in the last 30 days.", False, RGBColor(0, 0, 0), False)],
    [("When you signed up to join tonight\u2019s event, you became a member.", False, RGBColor(0, 0, 0), False)],
    [("Our Region", True, RGBColor(0, 0, 0), False), (" = South Central: 17 chapters, 7,800+ members.  Carolinas: 6 chapters, 560 members.", False, RGBColor(0, 0, 0), False)],
    [("The AI Collective Atlanta", True, RGBColor(0, 0, 0), False), (" = 4,800+ members, 115+ events since July 2025", False, RGBColor(0, 0, 0), False)],
    [("We meet 2 to 3 times a week all over the Atlanta area.", False, RGBColor(0, 0, 0), False)],
    [("24 organizers in Atlanta, and we\u2019re adding more to cover the metro area and lead committees and special interest groups.", False, RGBColor(0, 0, 0), False)],
], size=13.5, space_after=7)
place(body, t=1.25, h=4.2); body.text_frame.vertical_anchor = 1
from pptx.oxml.ns import qn
for i, para in enumerate(body.text_frame.paragraphs):
    pPr = para._p.get_or_add_pPr()
    if i in (0, 1, 4, 5):
        pPr.set("marL", "0"); pPr.set("indent", "0")
        for tag in ("a:buAutoNum", "a:buChar", "a:buNone"):
            for el in pPr.findall(qn(tag)): pPr.remove(el)
        pPr.append(pPr.makeelement(qn("a:buNone"), {}))
    if i in (1, 4, 5): para.space_before = Pt(8)

# ---- small cleanups: one font on the values slide, exercise times, Airbnb
for sl in p.slides:
    is_values = any(sh.has_text_frame and "COMMUNITY VALUES" in sh.text_frame.text for sh in sl.shapes)
    for sh in sl.shapes:
        if not sh.has_text_frame: continue
        for para in sh.text_frame.paragraphs:
            for r in para.runs:
                if is_values: r.font.name = "Calibri"
                r.text = r.text.replace("7:00 \u20137:45", "7:00 \u2013 7:45").replace("Air BnB", "Airbnb").replace("Next Vibe Night", "Next Vibe Nights")

# ---- slide order: the tools slide sits right before Build an App
def title_of(sl):
    return " | ".join(sh.text_frame.text.strip() for sh in sl.shapes if sh.has_text_frame)
lst = p.slides._sldIdLst; ids = list(lst); sl = list(p.slides)
ti = next(i for i, x in enumerate(sl) if "Four Tools to Vibe Code Tonight" in title_of(x))
el = ids[ti]; lst.remove(el)
sl2 = list(p.slides); bi2 = next(i for i, x in enumerate(sl2) if "Build an App" in title_of(x) and "Context is the most important part" in title_of(x))
lst.insert(bi2, el)

p.save(out)
print("saved", out, "slides", len(p.slides))

# ---- hype cycle slide: "you are here tonight" marker beside the peak (plain text, no shapes)
p2 = Presentation(out)
hype = next(s for s in p2.slides if any(sh.has_text_frame is False and sh.shape_type == 13 and sh.width > Inches(8) for sh in s.shapes) and len(s.shapes) == 1)
tb = hype.shapes.add_textbox(Inches(2.08), Inches(2.55), Inches(1.45), Inches(0.62))
tb.text_frame.word_wrap = True
r = tb.text_frame.paragraphs[0].add_run(); r.text = "← You are here tonight!"
r.font.size = Pt(15); r.font.bold = True; r.font.name = "Calibri"; r.font.color.rgb = RGBColor(0xF6, 0x73, 0x17)
p2.save(out)
print("hype marker added")
