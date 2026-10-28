import sys, zipfile, io
from pptx import Presentation
from PIL import Image
src,out=sys.argv[1],sys.argv[2]
p=Presentation(src); s=next(x for x in p.slides if any(sh.has_text_frame and sh.text_frame.text.strip()=="What's Next" for sh in x.shapes))
def setText(sh,txt):
    paras=sh.text_frame.paragraphs; runs=paras[0].runs
    runs[0].text=txt
    for r in runs[1:]: r.text=''
    for extra in paras[1:]:
        for r in extra.runs: r.text=''
named=[(sh.top,sh) for sh in s.shapes if sh.has_text_frame]
nxt=[sh for t,sh in named if sh.name=='Google Shape;339;p22'][0]
setText(nxt,'{{NEXT_VN}}')
evn=sorted([(t,sh) for t,sh in named if sh.name=='Google Shape;344;p22'],key=lambda x:x[0])
evd=sorted([(t,sh) for t,sh in named if sh.name=='Google Shape;345;p22'],key=lambda x:x[0])
for i,((t1,a),(t2,b)) in enumerate(zip(evn,evd),1):
    print('event',i,repr(a.text_frame.text),'|',repr(b.text_frame.text))
    setText(a,'{{EV%d_NAME}}'%i); setText(b,'{{EV%d_DETAILS}}'%i)
buf=io.BytesIO(); p.save(buf)
z=zipfile.ZipFile(buf); 
tiny=io.BytesIO(); Image.new('P',(1,1)).save(tiny,'GIF')
with zipfile.ZipFile(out,'w',zipfile.ZIP_DEFLATED) as o:
    for n in z.namelist():
        b=z.read(n)
        if n=='ppt/media/image2.gif': b=tiny.getvalue()
        o.writestr(n,b)
import os; print('template size',os.path.getsize(out))
